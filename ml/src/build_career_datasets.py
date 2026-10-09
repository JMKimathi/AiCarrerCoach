"""Build small, traceable career data tables from downloaded CSV and Excel exports.

Reads the LinkedIn postings/skills join and available O*NET occupation, skills,
interests, activities, and job-zone tables.
It does not retain job descriptions, company details, contact information, or
demographic fields. Job-market counts and O*NET occupation profiles are kept as
separate outputs because they answer different questions.

Usage:
  python ml/src/build_career_datasets.py
  python ml/src/build_career_datasets.py --input-dir "C:/Users/Mwiti/Downloads/Datasets"
"""

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path


ROOT_DIR = Path(__file__).resolve().parents[2]
DEFAULT_INPUT_DIR = Path.home() / "Downloads" / "Datasets"
DEFAULT_OUTPUT_DIR = ROOT_DIR / "ml" / "data" / "processed" / "derived"
OPTIONAL_ONET_ALIASES = {
    "career_interests": {"careerinteresttypes", "careerinterests", "interests"},
    "specific_interests": {"specificinterestareas", "specificinterests"},
    "work_activities": {"workactivities", "activities"},
    "job_zones": {"jobzones"},
    "occupation_data": {"occupationdata"},
    "job_titles": {"jobtitles"},
}


def open_csv(path: Path):
    return path.open("r", encoding="utf-8-sig", newline="", errors="replace")


def read_rows(path: Path):
    """Yield row dictionaries from a CSV or the first worksheet in an XLSX."""
    if path.suffix.casefold() == ".csv":
        with open_csv(path) as handle:
            yield from csv.DictReader(handle)
        return
    if path.suffix.casefold() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel inputs require openpyxl; run `pip install openpyxl`.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
            for values in rows:
                yield {header: ("" if value is None else str(value).strip())
                       for header, value in zip(headers, values) if header}
        finally:
            workbook.close()
        return
    raise ValueError(f"Unsupported input format: {path.name}")


def validate_columns(path: Path, required: set[str]) -> set[str]:
    with open_csv(path) as handle:
        reader = csv.DictReader(handle)
        available = set(reader.fieldnames or [])
    missing = required - available
    if missing:
        raise ValueError(f"{path.name} is missing required columns: {sorted(missing)}")
    return available


def clean_id(value: str | None) -> str:
    if value is None:
        return ""
    value = value.strip()
    return re.sub(r"\.0$", "", value)


def find_optional_onet_files(input_dir: Path) -> dict[str, Path]:
    found = {}
    for path in (*input_dir.glob("*.csv"), *input_dir.glob("*.xlsx")):
        normalized_name = re.sub(r"[^a-z]", "", path.stem.casefold())
        for category, aliases in OPTIONAL_ONET_ALIASES.items():
            if normalized_name in aliases:
                found[category] = path
    return found


def ensure_profile(profiles: dict[str, dict], code: str, title: str) -> dict:
    return profiles.setdefault(code, {
        "onet_soc_code": code,
        "occupation_title": title,
        "essential_skills": {},
        "software_skills": {},
        "career_interests": {},
        "specific_interests": {},
        "work_activities": {},
        "job_zone": "",
        "description": "",
        "job_titles": [],
    })


def build_onet_profiles(input_dir: Path, output_dir: Path, optional_files: dict[str, Path]) -> dict:
    essential_path = input_dir / "essential_skills.csv"
    software_path = input_dir / "software_skills.csv"
    profiles: dict[str, dict] = {}
    counts = {"essential_skill_rows": 0, "software_skill_rows": 0}

    with open_csv(essential_path) as handle:
        for row in csv.DictReader(handle):
            code = (row.get("O*NET-SOC Code") or "").strip()
            if not code:
                continue
            counts["essential_skill_rows"] += 1
            profile = ensure_profile(profiles, code, (row.get("Title") or "").strip())
            skill = (row.get("Element Name") or "").strip()
            if skill:
                scale = (row.get("Scale Name") or "").strip() or "unspecified"
                profile["essential_skills"].setdefault(skill, {})[scale] = (row.get("Data Value") or "").strip()

    with open_csv(software_path) as handle:
        for row in csv.DictReader(handle):
            code = (row.get("O*NET-SOC Code") or "").strip()
            if not code:
                continue
            counts["software_skill_rows"] += 1
            profile = ensure_profile(profiles, code, (row.get("Title") or "").strip())
            skill = (row.get("Element Name") or "").strip()
            if skill:
                example = (row.get("Workplace Example") or "").strip()
                entry_key = f"{skill} | {example}"
                profile["software_skills"][entry_key] = {
                    "name": skill,
                    "hot_technology": (row.get("Hot Technology") or "").strip(),
                    "in_demand": (row.get("In Demand") or "").strip(),
                    "workplace_example": example,
                }

    optional_rows = {}
    optional_requirements = {
        "career_interests": {"O*NET-SOC Code", "Title", "Element Name", "Scale Name", "Data Value"},
        "specific_interests": {"O*NET-SOC Code", "Title", "Element Name", "Scale Name", "Data Value"},
        "work_activities": {"O*NET-SOC Code", "Title", "Element Name", "Scale Name", "Data Value"},
        "job_zones": {"O*NET-SOC Code", "Title", "Job Zone"},
        "occupation_data": {"O*NET-SOC Code", "Title", "Description"},
        "job_titles": {"O*NET-SOC Code", "Title", "Job Title", "Short Title"},
    }
    for category, path in optional_files.items():
        if path.suffix.casefold() == ".csv":
            validate_columns(path, optional_requirements[category])
        optional_rows[category] = 0
        for row in read_rows(path):
            missing = optional_requirements[category] - set(row)
            if missing:
                raise ValueError(f"{path.name} is missing required columns: {sorted(missing)}")
            code = (row.get("O*NET-SOC Code") or "").strip()
            if not code:
                continue
            profile = ensure_profile(profiles, code, (row.get("Title") or "").strip())
            if category in ("career_interests", "specific_interests", "work_activities"):
                element = (row.get("Element Name") or "").strip()
                if element:
                    scale = (row.get("Scale Name") or "").strip() or "unspecified"
                    profile[category].setdefault(element, {})[scale] = (row.get("Data Value") or "").strip()
            elif category == "job_zones":
                profile["job_zone"] = (row.get("Job Zone") or "").strip()
            elif category == "occupation_data":
                profile["description"] = (row.get("Description") or "").strip()
            else:
                title = (row.get("Job Title") or row.get("Short Title") or "").strip()
                if title and title not in profile["job_titles"]:
                    profile["job_titles"].append(title)
            optional_rows[category] += 1

    output_path = output_dir / "onet_occupation_skills.jsonl"
    with output_path.open("w", encoding="utf-8", newline="\n") as handle:
        for code in sorted(profiles):
            profile = profiles[code]
            profile["essential_skills"] = [
                {"name": name, **values} for name, values in sorted(profile["essential_skills"].items())
            ]
            profile["software_skills"] = [
                values for _, values in sorted(profile["software_skills"].items())
            ]
            profile["career_interests"] = [
                {"name": name, **values} for name, values in sorted(profile["career_interests"].items())
            ]
            profile["specific_interests"] = [
                {"name": name, **values} for name, values in sorted(profile["specific_interests"].items())
            ]
            profile["work_activities"] = [
                {"name": name, **values} for name, values in sorted(profile["work_activities"].items())
            ]
            handle.write(json.dumps(profile, ensure_ascii=False) + "\n")

    counts["occupation_profiles"] = len(profiles)
    counts["optional_onet_rows"] = optional_rows
    counts["optional_onet_files_missing"] = sorted(set(OPTIONAL_ONET_ALIASES) - set(optional_files))
    counts["output"] = output_path.name
    return counts


def likely_kenya_location(location: str) -> bool:
    value = location.casefold()
    if "kenya" in value:
        return True
    cities = ("nairobi", "mombasa", "kisumu", "nakuru", "eldoret", "thika")
    return any(city in value for city in cities)


def build_linkedin_skill_demand(input_dir: Path, output_dir: Path) -> dict:
    skills_path = input_dir / "skills.csv"
    job_skills_path = input_dir / "job_skills.csv"
    postings_path = input_dir / "postings.csv"

    skill_names = {}
    with open_csv(skills_path) as handle:
        for row in csv.DictReader(handle):
            code = (row.get("skill_abr") or "").strip()
            name = (row.get("skill_name") or "").strip()
            if code and name:
                skill_names[code] = name

    skills_by_job: dict[str, set[str]] = defaultdict(set)
    unknown_skill_codes = 0
    with open_csv(job_skills_path) as handle:
        for row in csv.DictReader(handle):
            job_id = clean_id(row.get("job_id"))
            skill_code = (row.get("skill_abr") or "").strip()
            skill_name = skill_names.get(skill_code)
            if job_id and skill_name:
                skills_by_job[job_id].add(skill_name)
            elif job_id and skill_code:
                unknown_skill_codes += 1

    demand: Counter[tuple[str, str, str]] = Counter()
    stats = Counter()
    unknown_date_count = 0
    postings_with_skills = set()
    unique_posting_ids = set()
    locations = Counter()
    kenya_postings = 0

    with open_csv(postings_path) as handle:
        for row in csv.DictReader(handle):
            job_id = clean_id(row.get("job_id"))
            if not job_id:
                stats["rows_without_job_id"] += 1
                continue
            if job_id in unique_posting_ids:
                stats["duplicate_job_rows_skipped"] += 1
                continue
            unique_posting_ids.add(job_id)
            title = (row.get("title") or "").strip()
            location = (row.get("location") or "").strip() or "(unspecified)"
            locations[location] += 1
            if likely_kenya_location(location):
                kenya_postings += 1
            if not title:
                stats["rows_without_title"] += 1
            listed_time = (row.get("listed_time") or row.get("original_listed_time") or "").strip()
            try:
                stamp = float(listed_time)
                # LinkedIn exports normally express listed_time in milliseconds.
                seconds = stamp / 1000 if stamp > 10_000_000_000 else stamp
                year = str(datetime.fromtimestamp(seconds, tz=timezone.utc).year)
            except (ValueError, OverflowError, OSError):
                year = "(unknown)"
                unknown_date_count += 1
            work_type = (row.get("formatted_work_type") or row.get("work_type") or "").strip()

            job_skill_names = skills_by_job.get(job_id, set())
            if job_skill_names:
                postings_with_skills.add(job_id)
                for skill in job_skill_names:
                    demand[(location, title or "(unspecified)", skill)] += 1
            stats[f"posting_year:{year}"] += 1
            if work_type:
                stats[f"work_type:{work_type}"] += 1

    output_path = output_dir / "linkedin_job_skill_demand.csv"
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["location", "job_title", "skill", "posting_count"])
        for (location, title, skill), count in sorted(demand.items()):
            writer.writerow([location, title, skill, count])

    return {
        "postings_rows": sum(value for key, value in stats.items() if key.startswith("posting_year:")),
        "unique_posting_ids": len(unique_posting_ids),
        "postings_with_joined_skills": len(postings_with_skills),
        "unique_skill_codes": len(skill_names),
        "job_skill_join_ids": len(skills_by_job),
        "job_skill_rows_with_unknown_codes": unknown_skill_codes,
        "duplicate_job_rows_skipped": stats["duplicate_job_rows_skipped"],
        "distinct_locations": len(locations),
        "likely_kenya_postings": kenya_postings,
        "unknown_posting_dates": unknown_date_count,
        "posting_years": {key.split(":", 1)[1]: value for key, value in sorted(stats.items()) if key.startswith("posting_year:")},
        "work_types": {key.split(":", 1)[1]: value for key, value in sorted(stats.items()) if key.startswith("work_type:")},
        "top_locations": locations.most_common(20),
        "skill_demand_rows": len(demand),
        "output": output_path.name,
        "location_note": "Kenya count is a simple location-string heuristic (country name or common city names); review before using it as a market statistic.",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    required_files = (
        "essential_skills.csv", "software_skills.csv", "skills.csv",
        "job_skills.csv", "postings.csv",
    )
    missing = [name for name in required_files if not (args.input_dir / name).is_file()]
    if missing:
        parser.error(f"Missing required input files in {args.input_dir}: {', '.join(missing)}")

    required_columns = {
        "essential_skills.csv": {"O*NET-SOC Code", "Title", "Element Name", "Scale Name", "Data Value"},
        "software_skills.csv": {"O*NET-SOC Code", "Title", "Element Name", "Hot Technology", "In Demand", "Workplace Example"},
        "skills.csv": {"skill_abr", "skill_name"},
        "job_skills.csv": {"job_id", "skill_abr"},
        "postings.csv": {"job_id", "title", "location"},
    }
    headers = {name: validate_columns(args.input_dir / name, columns) for name, columns in required_columns.items()}
    if not ({"listed_time", "original_listed_time"} & headers["postings.csv"]):
        parser.error("postings.csv needs listed_time or original_listed_time to retain posting year.")
    optional_onet_files = find_optional_onet_files(args.input_dir)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    manifest = {
        "pipeline_version": "career-data-build-v1",
        "input_directory": str(args.input_dir),
        "output_directory": str(args.output_dir),
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_files": {
            name: {
                "size_bytes": (args.input_dir / name).stat().st_size,
                "modified_at": datetime.fromtimestamp((args.input_dir / name).stat().st_mtime, timezone.utc).isoformat(),
                "required_columns_validated": True,
            }
            for name in required_files
        },
        "optional_onet_files": {
            category: {"filename": path.name, "size_bytes": path.stat().st_size}
            for category, path in optional_onet_files.items()
        },
        "onet": build_onet_profiles(args.input_dir, args.output_dir, optional_onet_files),
        "linkedin": build_linkedin_skill_demand(args.input_dir, args.output_dir),
        "not_included": [
            "job_descriptions.csv (separate global dataset; not merged)",
            "company and employee tables",
            "benefit and salary tables",
            "student profiles or assessment answers",
        ],
        "interpretation": "LinkedIn counts describe skills mentioned in postings. O*NET profiles describe U.S. occupational skill requirements, interests, activities, and preparation zones. Neither is a labeled student career-success dataset.",
    }
    manifest_path = args.output_dir / "career_data_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
